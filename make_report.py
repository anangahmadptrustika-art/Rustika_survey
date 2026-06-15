#!/usr/bin/env python3
"""
make_report.py — Laporan survey kondisi jalan (Fase 4): PDF + Excel A4.

Dari detections.csv satu sesi survey, menghasilkan laporan profesional siap kirim:
  • Ringkasan: jumlah kerusakan per jenis + tingkat keparahan (dari ukuran bbox)
  • Tabel detail tiap kerusakan (lokasi GPS, keparahan, confidence, waktu, foto)
  • Peta/sebaran titik kerusakan
  • Lampiran foto

PENTING — penghitungan jujur:
  CSV mencatat deteksi PER-FRAME, jadi satu lubang muncul di banyak frame. Skrip
  ini meng-KLASTER deteksi menjadi "kerusakan unik" (berdasarkan kedekatan GPS,
  atau kedekatan antar-frame bila tanpa GPS) sebelum dihitung. Tanpa ini, laporan
  akan melebih-lebihkan jumlah kerusakan.

Keparahan (Ringan/Sedang/Berat) ditaksir dari rasio luas bbox terhadap frame.
CATATAN: ini proksi kasar — jarak kamera / ketinggian drone memengaruhi ukuran
bbox, dan retak yang tipis-memanjang bisa "terlihat kecil". Kalibrasi ambang
(--sev-low/--sev-high) dengan footage sendiri, dan verifikasi manual untuk klaim
formal ke klien.

Contoh (sekali perintah):
  python make_report.py output/survey_20260615_120000/detections.csv \
      --road "Ruas Malili–Wawondula" --surveyor "Anang A." --date 2026-06-15
"""
from __future__ import annotations

import argparse
import csv
import math
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

SEVERITY = ["Ringan", "Sedang", "Berat"]
SEV_HEX = {"Ringan": "#2ca02c", "Sedang": "#ff7f0e", "Berat": "#d62728"}
CLASS_ID_LABEL = {
    "longitudinal_crack": "Retak Memanjang",
    "transverse_crack": "Retak Melintang",
    "alligator_crack": "Retak Buaya",
    "pothole": "Lubang",
}


def label_id(cls: str) -> str:
    return CLASS_ID_LABEL.get(cls, cls)


def haversine_m(lat1, lon1, lat2, lon2) -> float:
    R = 6371000.0
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp = math.radians(lat2 - lat1)
    dl = math.radians(lon2 - lon1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * R * math.asin(math.sqrt(a))


def severity_for(rel_area: float, low: float, high: float) -> str:
    if rel_area >= high:
        return "Berat"
    if rel_area >= low:
        return "Sedang"
    return "Ringan"


# --------------------------------------------------------------------------- #
# Baca & klaster
# --------------------------------------------------------------------------- #
def load_rows(csv_path: Path) -> list[dict]:
    rows = []
    with open(csv_path, newline="", encoding="utf-8") as f:
        for d in csv.DictReader(f):
            try:
                fw = float(d.get("frame_w") or 0)
                fh = float(d.get("frame_h") or 0)
                area = float(d.get("area_px") or 0)
            except ValueError:
                continue
            lat = lon = None
            if d.get("lat") not in ("", None) and d.get("lon") not in ("", None):
                try:
                    lat = float(d["lat"]); lon = float(d["lon"])
                except ValueError:
                    pass
            area_m2 = None
            if d.get("area_m2") not in ("", None):
                try:
                    area_m2 = float(d["area_m2"])
                except ValueError:
                    pass
            rows.append({
                "class": d.get("class_name", "?"),
                "conf": float(d.get("confidence") or 0),
                "frame": int(float(d.get("frame_idx") or 0)),
                "epoch": float(d.get("epoch_s") or 0),
                "time": d.get("timestamp_iso", ""),
                "rel_area": (area / (fw * fh)) if fw and fh else 0.0,
                "area_m2": area_m2,
                "lat": lat, "lon": lon,
                "screenshot": d.get("screenshot", ""),
            })
    return rows


def cluster_detections(rows, merge_dist_m, merge_frames):
    """Gabungkan deteksi per-frame menjadi kerusakan unik."""
    has_gps = any(r["lat"] is not None for r in rows)
    rows = sorted(rows, key=lambda r: (r["epoch"], r["frame"]))
    clusters = []
    for r in rows:
        match = None
        for c in reversed(clusters[-80:]):       # cek klaster terbaru saja (cepat)
            if c["class"] != r["class"]:
                continue
            if has_gps and r["lat"] is not None and c["lat"] is not None:
                if haversine_m(c["lat"], c["lon"], r["lat"], r["lon"]) <= merge_dist_m:
                    match = c; break
            elif abs(r["frame"] - c["last_frame"]) <= merge_frames:
                match = c; break
        if match is None:
            clusters.append({
                "class": r["class"], "max_conf": r["conf"],
                "max_rel_area": r["rel_area"], "max_area_m2": r["area_m2"],
                "n_det": 1,
                "first_frame": r["frame"], "last_frame": r["frame"],
                "first_time": r["time"], "lat": r["lat"], "lon": r["lon"],
                "_lat_sum": r["lat"] or 0.0, "_lon_sum": r["lon"] or 0.0,
                "_n_coord": 1 if r["lat"] is not None else 0,
                "screenshot": r["screenshot"],
            })
        else:
            match["n_det"] += 1
            match["max_conf"] = max(match["max_conf"], r["conf"])
            match["max_rel_area"] = max(match["max_rel_area"], r["rel_area"])
            if r["area_m2"] is not None:
                match["max_area_m2"] = max(match["max_area_m2"] or 0.0, r["area_m2"])
            match["last_frame"] = r["frame"]
            if r["lat"] is not None:
                match["_lat_sum"] += r["lat"]; match["_lon_sum"] += r["lon"]
                match["_n_coord"] += 1
                match["lat"] = match["_lat_sum"] / match["_n_coord"]
                match["lon"] = match["_lon_sum"] / match["_n_coord"]
            if not match["screenshot"] and r["screenshot"]:
                match["screenshot"] = r["screenshot"]
    return clusters, has_gps


def summarize(clusters, low, high, low_m2, high_m2):
    for c in clusters:
        if c.get("max_area_m2") is not None:
            # Keparahan dari luas nyata (m²) — lebih bermakna bila terkalibrasi.
            c["severity"] = severity_for(c["max_area_m2"], low_m2, high_m2)
        else:
            c["severity"] = severity_for(c["max_rel_area"], low, high)
    per_class = Counter(c["class"] for c in clusters)
    per_sev = Counter(c["severity"] for c in clusters)
    matrix = defaultdict(lambda: Counter())
    for c in clusters:
        matrix[c["class"]][c["severity"]] += 1
    return per_class, per_sev, matrix


# --------------------------------------------------------------------------- #
# Grafik (offline, matplotlib Agg)
# --------------------------------------------------------------------------- #
def build_charts(clusters, per_class, matrix, out_dir: Path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    paths = {}
    # 1. Bar bertumpuk: jumlah per jenis x keparahan.
    classes = list(per_class.keys())
    if classes:
        fig, ax = plt.subplots(figsize=(6.5, 3.2))
        bottom = [0] * len(classes)
        for sev in SEVERITY:
            vals = [matrix[c][sev] for c in classes]
            ax.bar([label_id(c) for c in classes], vals, bottom=bottom,
                   label=sev, color=SEV_HEX[sev])
            bottom = [b + v for b, v in zip(bottom, vals)]
        ax.set_ylabel("Jumlah kerusakan")
        ax.set_title("Kerusakan per Jenis & Tingkat Keparahan")
        ax.legend(title="Keparahan", fontsize=8)
        plt.xticks(rotation=15, ha="right", fontsize=8)
        plt.tight_layout()
        p = out_dir / "_chart_severity.png"
        fig.savefig(p, dpi=130); plt.close(fig)
        paths["severity"] = p

    # 2. Sebaran titik (lon vs lat), offline (tanpa tile).
    pts = [(c["lon"], c["lat"], c["class"]) for c in clusters if c["lat"] is not None]
    if pts:
        fig, ax = plt.subplots(figsize=(6.5, 4.0))
        from make_map import CLASS_COLORS, DEFAULT_COLOR
        for cls in {p[2] for p in pts}:
            xs = [p[0] for p in pts if p[2] == cls]
            ys = [p[1] for p in pts if p[2] == cls]
            ax.scatter(xs, ys, s=30, label=label_id(cls),
                       color=CLASS_COLORS.get(cls, DEFAULT_COLOR), edgecolors="k",
                       linewidths=0.3)
        ax.set_xlabel("Longitude"); ax.set_ylabel("Latitude")
        ax.set_title("Sebaran Lokasi Kerusakan")
        ax.legend(fontsize=8); ax.ticklabel_format(useOffset=False, style="plain")
        ax.grid(True, alpha=0.3)
        plt.tight_layout()
        p = out_dir / "_chart_map.png"
        fig.savefig(p, dpi=130); plt.close(fig)
        paths["map"] = p
    return paths


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def build_excel(clusters, meta, per_class, matrix, has_size, out_path: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    bold = Font(bold=True)
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    hdr_font = Font(bold=True, color="FFFFFF")
    thin = Border(*[Side(style="thin", color="DDDDDD")] * 4)

    # Sheet 1: Ringkasan
    ws = wb.active; ws.title = "Ringkasan"
    ws["A1"] = "LAPORAN SURVEY KONDISI JALAN"; ws["A1"].font = Font(bold=True, size=14)
    r = 3
    for k, v in meta.items():
        ws.cell(r, 1, k).font = bold; ws.cell(r, 2, v); r += 1
    r += 1
    ws.cell(r, 1, "Total kerusakan unik").font = bold
    ws.cell(r, 2, len(clusters)); r += 2

    ws.cell(r, 1, "Jenis").font = hdr_font; ws.cell(r, 1).fill = hdr_fill
    for j, sev in enumerate(SEVERITY + ["Total"]):
        cc = ws.cell(r, 2 + j, sev); cc.font = hdr_font; cc.fill = hdr_fill
    r += 1
    for cls in per_class:
        ws.cell(r, 1, label_id(cls))
        tot = 0
        for j, sev in enumerate(SEVERITY):
            v = matrix[cls][sev]; tot += v
            ws.cell(r, 2 + j, v)
        ws.cell(r, 5, tot).font = bold
        r += 1
    for col, w in {"A": 26, "B": 16, "C": 12, "D": 12, "E": 12}.items():
        ws.column_dimensions[col].width = w

    # Sheet 2: Detail
    ws2 = wb.create_sheet("Detail Kerusakan")
    headers = ["No", "Jenis", "Keparahan", "Latitude", "Longitude",
               "Confidence", "Luas bbox (%)", "Luas (m²)", "Frame", "Waktu", "Foto"]
    for j, h in enumerate(headers, 1):
        c = ws2.cell(1, j, h); c.font = hdr_font; c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center")
    for i, c in enumerate(sorted(clusters, key=lambda x: (x["class"], -x["max_rel_area"])), 1):
        row = [
            i, label_id(c["class"]), c["severity"],
            round(c["lat"], 7) if c["lat"] is not None else "",
            round(c["lon"], 7) if c["lon"] is not None else "",
            round(c["max_conf"], 3), round(c["max_rel_area"] * 100, 3),
            round(c["max_area_m2"], 4) if c.get("max_area_m2") is not None else "",
            c["first_frame"], c["first_time"], c["screenshot"],
        ]
        for j, v in enumerate(row, 1):
            cell = ws2.cell(i + 1, j, v); cell.border = thin
        sev_cell = ws2.cell(i + 1, 3)
        sev_cell.fill = PatternFill("solid", fgColor=SEV_HEX[c["severity"]].lstrip("#"))
        sev_cell.font = Font(color="FFFFFF", bold=True)
    widths = [5, 18, 11, 13, 13, 11, 14, 11, 8, 24, 26]
    for j, w in enumerate(widths, 1):
        ws2.column_dimensions[chr(64 + j)].width = w
    ws2.freeze_panes = "A2"

    wb.save(out_path)
    return out_path


# --------------------------------------------------------------------------- #
# PDF (A4)
# --------------------------------------------------------------------------- #
def build_pdf(clusters, meta, per_class, per_sev, matrix, charts,
              shots_dir: Path, has_size: bool, out_path: Path, max_photos: int):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                    TableStyle, Image as RLImage)

    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Title"], fontSize=18, spaceAfter=2)
    sub = ParagraphStyle("sub", parent=styles["Normal"], fontSize=10,
                         textColor=colors.HexColor("#555555"), spaceAfter=10)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=13,
                        textColor=colors.HexColor("#1F4E78"), spaceBefore=10)
    small = ParagraphStyle("small", parent=styles["Normal"], fontSize=8,
                           textColor=colors.HexColor("#777777"))

    doc = SimpleDocTemplate(str(out_path), pagesize=A4,
                            topMargin=18 * mm, bottomMargin=16 * mm,
                            leftMargin=16 * mm, rightMargin=16 * mm,
                            title="Laporan Survey Kondisi Jalan")
    story = []
    story.append(Paragraph("LAPORAN SURVEY KONDISI JALAN", h1))
    story.append(Paragraph("Rustika Citra Group — Deteksi Kerusakan Jalan (AI)", sub))

    meta_rows = [[k, str(v)] for k, v in meta.items()]
    meta_rows.append(["Total kerusakan unik", str(len(clusters))])
    mt = Table(meta_rows, colWidths=[55 * mm, 110 * mm])
    mt.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#1F4E78")),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#EEEEEE")),
    ]))
    story.append(mt)

    # Ringkasan
    story.append(Paragraph("Ringkasan", h2))
    head = ["Jenis"] + SEVERITY + ["Total"]
    data = [head]
    for cls in per_class:
        tot = sum(matrix[cls].values())
        data.append([label_id(cls)] + [str(matrix[cls][s]) for s in SEVERITY] + [str(tot)])
    data.append(["TOTAL"] + [str(per_sev.get(s, 0)) for s in SEVERITY] + [str(len(clusters))])
    st = Table(data, colWidths=[55 * mm] + [30 * mm] * 4)
    sty = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#EEF3F8")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CCCCCC")),
        ("FONTSIZE", (0, 0), (-1, -1), 9), ("ALIGN", (1, 0), (-1, -1), "CENTER"),
    ]
    # warnai header keparahan
    for j, s in enumerate(SEVERITY, 1):
        sty.append(("TEXTCOLOR", (j, 0), (j, 0), colors.white))
    st.setStyle(TableStyle(sty))
    story.append(st)

    if "severity" in charts:
        story.append(Spacer(1, 6))
        story.append(RLImage(str(charts["severity"]), width=160 * mm, height=78 * mm))

    if "map" in charts:
        story.append(Paragraph("Sebaran Lokasi", h2))
        story.append(RLImage(str(charts["map"]), width=160 * mm, height=98 * mm))
        story.append(Paragraph(
            "Catatan: sebaran berbasis koordinat GPS. Peta jalan interaktif ada di "
            "map.html; untuk basemap offline gunakan detections.geojson di QGIS.", small))

    # Tabel detail (maksimal beberapa, sisanya di Excel)
    story.append(Paragraph("Detail Kerusakan", h2))
    size_col = "Luas m²" if has_size else "Luas%"
    det_head = ["No", "Jenis", "Keparahan", "Lat", "Lon", "Conf", size_col]
    det = [det_head]
    ordered = sorted(clusters, key=lambda x: (x["class"], -x["max_rel_area"]))
    for i, c in enumerate(ordered[:40], 1):
        size_val = (f'{c["max_area_m2"]:.3f}' if has_size and c.get("max_area_m2")
                    is not None else f'{c["max_rel_area"]*100:.2f}')
        det.append([
            str(i), label_id(c["class"]), c["severity"],
            f'{c["lat"]:.6f}' if c["lat"] is not None else "-",
            f'{c["lon"]:.6f}' if c["lon"] is not None else "-",
            f'{c["max_conf"]:.2f}', size_val,
        ])
    dt = Table(det, colWidths=[10*mm, 32*mm, 22*mm, 28*mm, 28*mm, 16*mm, 18*mm], repeatRows=1)
    dsty = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#DDDDDD")),
        ("FONTSIZE", (0, 0), (-1, -1), 8), ("ALIGN", (0, 0), (0, -1), "CENTER"),
        ("ALIGN", (5, 0), (-1, -1), "CENTER"),
    ]
    for i, c in enumerate(ordered[:40], 1):
        dsty.append(("TEXTCOLOR", (2, i), (2, i), colors.HexColor(SEV_HEX[c["severity"]])))
        dsty.append(("FONTNAME", (2, i), (2, i), "Helvetica-Bold"))
    dt.setStyle(TableStyle(dsty))
    story.append(dt)
    if len(ordered) > 40:
        story.append(Paragraph(f"... {len(ordered)-40} kerusakan lain ada di file Excel.", small))

    # Lampiran foto
    photos = [(c, shots_dir / c["screenshot"]) for c in ordered
              if c["screenshot"] and (shots_dir / c["screenshot"]).exists()][:max_photos]
    if photos:
        story.append(Paragraph("Lampiran Foto", h2))
        cells = []
        for c, pth in photos:
            cap = Paragraph(f'{label_id(c["class"])} — <b>{c["severity"]}</b>', small)
            cells.append([RLImage(str(pth), width=52 * mm, height=39 * mm), cap])
        # susun grid 3 kolom
        grid = []
        rowbuf = []
        for cell in cells:
            inner = Table([[cell[0]], [cell[1]]])
            inner.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "CENTER")]))
            rowbuf.append(inner)
            if len(rowbuf) == 3:
                grid.append(rowbuf); rowbuf = []
        if rowbuf:
            while len(rowbuf) < 3:
                rowbuf.append("")
            grid.append(rowbuf)
        g = Table(grid, colWidths=[56 * mm] * 3)
        g.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
        story.append(g)

    story.append(Spacer(1, 8))
    story.append(Paragraph(
        "Dihasilkan otomatis oleh Road Survey AI. Keparahan ditaksir dari ukuran "
        "bbox (proksi) — verifikasi manual disarankan untuk klaim formal.", small))

    doc.build(story)
    return out_path


# --------------------------------------------------------------------------- #
# Orkestrasi
# --------------------------------------------------------------------------- #
def build_report(csv_path, out_dir, meta, want_pdf=True, want_excel=True,
                 merge_dist=8.0, merge_frames=20, sev_low=0.003, sev_high=0.015,
                 sev_low_m2=0.05, sev_high_m2=0.25, max_photos=9):
    csv_path = Path(csv_path); out_dir = Path(out_dir)
    rows = load_rows(csv_path)
    if not rows:
        raise SystemExit("CSV kosong / tidak ada deteksi.")
    clusters, has_gps = cluster_detections(rows, merge_dist, merge_frames)
    per_class, per_sev, matrix = summarize(clusters, sev_low, sev_high,
                                           sev_low_m2, sev_high_m2)
    has_size = any(c.get("max_area_m2") is not None for c in clusters)
    shots_dir = out_dir / "screenshots"

    print(f"  Deteksi mentah     : {len(rows)}")
    print(f"  Kerusakan unik     : {len(clusters)}  "
          f"({'klaster GPS' if has_gps else 'klaster antar-frame'})")
    print(f"  Keparahan dari     : {'luas m² (kalibrasi)' if has_size else 'ukuran bbox relatif'}")
    for cls, n in per_class.most_common():
        print(f"     - {label_id(cls):<18}: {n}")

    charts = build_charts(clusters, per_class, matrix, out_dir)
    outputs = []
    if want_excel:
        xlsx = build_excel(clusters, meta, per_class, matrix, has_size,
                           out_dir / "laporan.xlsx")
        outputs.append(xlsx); print(f"  Excel : {xlsx}")
    if want_pdf:
        pdf = build_pdf(clusters, meta, per_class, per_sev, matrix, charts,
                        shots_dir, has_size, out_dir / "laporan.pdf", max_photos)
        outputs.append(pdf); print(f"  PDF   : {pdf}")
    # bersihkan chart sementara
    for p in charts.values():
        try:
            Path(p).unlink()
        except OSError:
            pass
    return outputs


def parse_args(argv=None):
    p = argparse.ArgumentParser(description="Laporan survey kondisi jalan (PDF+Excel)")
    p.add_argument("csv", help="Path detections.csv")
    p.add_argument("--road", default=None, help="Nama ruas/jalan yang disurvey")
    p.add_argument("--surveyor", default="-", help="Nama surveyor")
    p.add_argument("--date", default=None, help="Tanggal survey (default: hari ini)")
    p.add_argument("--out", default=None, help="Folder output (default: folder CSV)")
    p.add_argument("--pdf-only", action="store_true")
    p.add_argument("--excel-only", action="store_true")
    p.add_argument("--merge-dist", type=float, default=8.0,
                   help="Jarak (m) menggabungkan deteksi jadi 1 kerusakan (mode GPS).")
    p.add_argument("--merge-frames", type=int, default=20,
                   help="Jeda frame menggabungkan deteksi (mode tanpa GPS).")
    p.add_argument("--sev-low", type=float, default=0.003,
                   help="Ambang luas bbox (fraksi) Ringan->Sedang (tanpa kalibrasi).")
    p.add_argument("--sev-high", type=float, default=0.015,
                   help="Ambang luas bbox (fraksi) Sedang->Berat (tanpa kalibrasi).")
    p.add_argument("--sev-low-m2", type=float, default=0.05,
                   help="Ambang luas m² Ringan->Sedang (bila terkalibrasi).")
    p.add_argument("--sev-high-m2", type=float, default=0.25,
                   help="Ambang luas m² Sedang->Berat (bila terkalibrasi).")
    p.add_argument("--max-photos", type=int, default=9)
    return p.parse_args(argv)


def main(argv=None) -> int:
    args = parse_args(argv)
    csv_path = Path(args.csv)
    if not csv_path.exists():
        raise SystemExit(f"CSV tidak ditemukan: {csv_path}")
    out_dir = Path(args.out) if args.out else csv_path.parent
    meta = {
        "Ruas jalan": args.road or csv_path.parent.name,
        "Tanggal survey": args.date or datetime.now().strftime("%Y-%m-%d"),
        "Surveyor": args.surveyor,
        "Konsultan": "Rustika Citra Group",
        "Wilayah": "Kabupaten Luwu Timur",
    }
    build_report(csv_path, out_dir, meta,
                 want_pdf=not args.excel_only, want_excel=not args.pdf_only,
                 merge_dist=args.merge_dist, merge_frames=args.merge_frames,
                 sev_low=args.sev_low, sev_high=args.sev_high,
                 sev_low_m2=args.sev_low_m2, sev_high_m2=args.sev_high_m2,
                 max_photos=args.max_photos)
    print("  Selesai. Laporan siap kirim.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
